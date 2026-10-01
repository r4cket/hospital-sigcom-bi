# -*- coding: utf-8 -*-
"""
Carga un Cubo 9 SIGCOM (Dashboard_SIGCOM.xlsx, hoja 'Base_Insumos') a Postgres.

Reaplica el esquema (/schema/01_schema.sql) y reemplaza por completo los datos
de centros_costo / costos_mensuales / produccion_mensual. Controla la cuadratura
RH + GG + Insumos + Indirecto == total del Excel en cada fila.

  python load_sigcom.py --xlsx /incoming/Dashboard_SIGCOM.xlsx [--strict]

Conexion por variables de entorno:
  PG_HOST PG_PORT PG_DATABASE PG_USER PG_PASSWORD  (PGSSLMODE opcional)

--strict  -> devuelve codigo 3 si hay avisos de cuadratura (no cambia la carga).
"""
import argparse
import os
import subprocess
import sys
import tempfile
import unicodedata
from collections import Counter

import openpyxl

MESES = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN",
         "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]


def slug(nombre):
    s = unicodedata.normalize("NFKD", nombre).encode("ascii", "ignore").decode()
    s = "".join(c if c.isalnum() else "_" for c in s).upper()
    while "__" in s:
        s = s.replace("__", "_")
    return s.strip("_")[:40]


def tipo_for(nombre, clasif):
    n = unicodedata.normalize("NFKD", nombre).encode("ascii", "ignore").decode().upper()
    if n.startswith("HOSPITALIZACION ") and "EN CASA" not in n:
        return "hospitalizacion"
    if "TRATAMIENTO INTENSIVO" in n:
        return "uti"
    if clasif.get(nombre) == "Apoyo":
        return "apoyo"
    return "ambulatorio"


def num(v):
    if v is None:
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def sql_str(s):
    return "'" + s.replace("'", "''") + "'"


def build_data_sql(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)

    if "Base_Insumos" not in wb.sheetnames:
        raise SystemExit(f"ERROR: el Excel no tiene la hoja 'Base_Insumos' (hojas: {wb.sheetnames})")

    clasif = {}
    if "Ranking_CC" in wb.sheetnames:
        for r in wb["Ranking_CC"].iter_rows(values_only=True):
            if r[1] and r[2] and str(r[2]).strip() in ("Final", "Apoyo"):
                clasif[str(r[1]).strip()] = str(r[2]).strip()

    rows = list(wb["Base_Insumos"].iter_rows(values_only=True))
    hdr = {name: i for i, name in enumerate(rows[0])}
    for req in ("CC", "Año"):
        if req not in hdr:
            raise SystemExit(f"ERROR: falta la columna '{req}' en Base_Insumos")

    centros = {}          # nombre -> (codigo, tipo, clasificacion)
    costos = []           # (codigo, 'YYYY-MM-01', rh, gg, ins, ind)
    prod = []             # (codigo, 'YYYY-MM-01', egresos, dco, camas)
    skipped = 0
    warn = []

    for r in rows[1:]:
        if r[hdr["CC"]] is None:
            continue
        nombre = str(r[hdr["CC"]]).strip()
        anio = int(r[hdr["Año"]])
        codigo = slug(nombre)
        if nombre not in centros:
            centros[nombre] = (codigo, tipo_for(nombre, clasif), clasif.get(nombre, "Final"))

        for mi, mon in enumerate(MESES, start=1):
            rh = num(r[hdr[f"rh_{mon}"]])
            gg = num(r[hdr[f"gg_{mon}"]])
            ins = num(r[hdr[f"ins_total_{mon}"]])
            ind = num(r[hdr[f"indirecto_{mon}"]])
            tot = num(r[hdr[f"total_{mon}"]])
            egr = num(r[hdr[f"produccion_{mon}"]])
            dco = num(r[hdr[f"dco_{mon}"]])
            cam = num(r[hdr[f"camas_{mon}"]])

            if rh == 0 and gg == 0 and ins == 0 and ind == 0 and tot == 0 and egr == 0 and dco == 0:
                skipped += 1
                continue

            periodo = f"{anio:04d}-{mi:02d}-01"
            costos.append((codigo, periodo, rh, gg, ins, ind))
            prod.append((codigo, periodo, round(egr), round(dco), round(cam)))

            calc = rh + gg + ins + ind
            if tot and abs(calc - tot) > 1.0:
                warn.append(f"{nombre} {periodo}: RH+GG+INS+IND={calc:,.0f} vs total Excel={tot:,.0f} (dif {calc - tot:,.0f})")

    if not costos:
        raise SystemExit("ERROR: no se encontro ninguna fila con actividad en Base_Insumos")

    lines = []
    lines.append("-- Generado por load_sigcom.py")
    lines.append(f"-- {len(centros)} centros | {len(costos)} costos_mensuales | {len(prod)} produccion_mensual")
    lines.append("BEGIN;")
    lines.append("")
    lines.append("INSERT INTO centros_costo (codigo, nombre, tipo, clasificacion) VALUES")
    vals = [f"    ({sql_str(cod)}, {sql_str(nom)}, {sql_str(tp)}, {sql_str(cl)})"
            for nom, (cod, tp, cl) in sorted(centros.items(), key=lambda x: x[1][0])]
    lines.append(",\n".join(vals) + ";")
    lines.append("")
    lines.append("INSERT INTO costos_mensuales (centro_costo_id, periodo, recurso_humano, gastos_generales, insumos, costo_indirecto) VALUES")
    vals = [f"    ((SELECT id FROM centros_costo WHERE codigo={sql_str(cod)}), DATE '{per}', {rh:.2f}, {gg:.2f}, {ins:.2f}, {ind:.2f})"
            for cod, per, rh, gg, ins, ind in costos]
    lines.append(",\n".join(vals) + ";")
    lines.append("")
    lines.append("INSERT INTO produccion_mensual (centro_costo_id, periodo, egresos, dias_cama_ocupados, camas_dotacion) VALUES")
    vals = [f"    ((SELECT id FROM centros_costo WHERE codigo={sql_str(cod)}), DATE '{per}', {egr}, {dco}, {cam})"
            for cod, per, egr, dco, cam in prod]
    lines.append(",\n".join(vals) + ";")
    lines.append("")
    lines.append("COMMIT;")

    stats = {
        "centros": len(centros),
        "costos": len(costos),
        "prod": len(prod),
        "skipped": skipped,
        "per_min": min(p[1] for p in costos),
        "per_max": max(p[1] for p in costos),
        "tipos": Counter(t for _, (_, t, _) in centros.items()),
    }
    return "\n".join(lines), stats, warn


def psql(args, sql_file):
    env = dict(os.environ)
    env["PGPASSWORD"] = env.get("PG_PASSWORD", "")
    cmd = [
        "psql",
        "-h", env.get("PG_HOST", "postgres"),
        "-p", env.get("PG_PORT", "5432"),
        "-U", env.get("PG_USER", "postgres"),
        "-d", env.get("PG_DATABASE", "postgres"),
        "-v", "ON_ERROR_STOP=1",
        "-q", "-f", sql_file,
    ]
    subprocess.run(cmd, check=True, env=env)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--schema", default="/schema/01_schema.sql")
    ap.add_argument("--strict", action="store_true",
                    default=os.environ.get("STRICT", "0") == "1")
    a = ap.parse_args()

    if not os.path.exists(a.xlsx):
        raise SystemExit(f"ERROR: no existe {a.xlsx}")

    print(f"[load] leyendo {a.xlsx}", flush=True)
    data_sql, stats, warn = build_data_sql(a.xlsx)

    with tempfile.NamedTemporaryFile("w", suffix=".sql", delete=False, encoding="utf-8") as fh:
        fh.write(data_sql)
        data_path = fh.name

    print(f"[load] aplicando esquema {a.schema}", flush=True)
    psql(a, a.schema)
    print("[load] cargando datos", flush=True)
    psql(a, data_path)
    os.unlink(data_path)

    print(f"[load] OK  centros={stats['centros']}  costos={stats['costos']}  "
          f"produccion={stats['prod']}  periodos {stats['per_min']}..{stats['per_max']}  "
          f"(celdas-mes sin actividad omitidas: {stats['skipped']})", flush=True)
    print(f"[load] tipos: {dict(stats['tipos'])}", flush=True)

    if warn:
        print(f"[load] !! {len(warn)} avisos de cuadratura (dif > $1):", flush=True)
        for w in warn[:15]:
            print("   " + w, flush=True)
        if a.strict:
            print("[load] --strict activo -> saliendo con codigo 3", flush=True)
            sys.exit(3)
    else:
        print("[load] cuadratura OK en todas las filas", flush=True)


if __name__ == "__main__":
    main()
